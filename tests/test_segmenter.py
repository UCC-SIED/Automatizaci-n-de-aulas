# -*- coding: utf-8 -*-
"""Tests de la segmentación de DOCX por sección."""

import docx
import pytest

from maquetador.extract.segmenter import segmentar_docx


def _crear_docx(tmp_path, parrafos):
    doc = docx.Document()
    for texto in parrafos:
        doc.add_paragraph(texto)
    ruta = tmp_path / "modulo.docx"
    doc.save(ruta)
    return ruta


def test_subtitulo_objetivo_dentro_de_seccion_no_secuestra_la_seccion(tmp_path):
    """Un sub-título 'Objetivo financiero' que abre una página numerada NO debe
    confundirse con el marcador especial 'Objetivos' del módulo: el cuerpo de la
    sección tiene que quedar dentro de la página, no dentro del bloque objetivos."""
    ruta = _crear_docx(tmp_path, [
        "Introducción",
        "Texto de introducción del módulo.",
        "Objetivos",
        "Que el estudiante comprenda los instrumentos financieros.",
        "1.1. Objetivo financiero. Principios financieros",
        "Objetivo financiero",
        "El cuerpo real de la sección uno punto uno va aquí.",
        "1.2. Otra sección",
        "Cuerpo de la sección uno punto dos.",
    ])
    marcadores = {
        "item_1": "1.1. Objetivo financiero. Principios financieros",
        "item_2": "1.2. Otra sección",
    }

    secciones, _img, faltantes, _com, _otp = segmentar_docx(ruta, marcadores)

    assert "item_1" not in faltantes
    assert "El cuerpo real de la sección uno punto uno" in secciones.get("item_1", "")
    # El cuerpo de 1.1 no debe haberse filtrado al bloque de objetivos.
    assert "cuerpo real de la sección uno punto uno" not in \
        secciones.get("objetivos", "").lower()


def test_encabezado_con_nota_al_pie_igual_matchea(tmp_path):
    """Un encabezado de sección con una nota al pie pegada ('… Trabajo Final [26]')
    debe cortarse igual: la referencia [N] que mete mammoth no está en el título de
    la planilla y no debe impedir el match (regresión: la sección quedaba vacía y su
    contenido se lo tragaba la sección anterior)."""
    ruta = _crear_docx(tmp_path, [
        "1.6. Conocer la evaluación",
        "Cuerpo de la sección uno punto seis.",
        "1.7. Comprender el papel de la IA [26]",
        "El cuerpo real de la sección uno punto siete va aquí.",
        "1.8. Guía para comenzar",
        "Cuerpo de la sección uno punto ocho.",
    ])
    marcadores = {
        "item_6": "1.6. Conocer la evaluación",
        "item_7": "1.7. Comprender el papel de la IA",   # sin la nota [26]
        "item_8": "1.8. Guía para comenzar",
    }

    secciones, _img, faltantes, _com, _otp = segmentar_docx(ruta, marcadores)

    assert "item_7" not in faltantes
    assert "cuerpo real de la sección uno punto siete" in secciones.get("item_7", "").lower()
    # No debe haberse filtrado a la sección anterior.
    assert "cuerpo real de la sección uno punto siete" not in secciones.get("item_6", "").lower()


def test_subtitulo_estilo_subtitle_se_convierte_en_h3(tmp_path):
    """Mammoth solo mapea por defecto "Heading 1".."Heading 6" a <hN>: el
    estilo Word "Subtitle" que los asesores usan para el subtítulo debajo del
    título de sección quedaba como <p> suelto (sin negrita ni marca alguna),
    invisible para el resto del pipeline. Regresión: "Planificación,
    aseguramiento y control de la calidad" no salía como <h3>."""
    doc = docx.Document()
    doc.add_paragraph("1.4. Gestión de la calidad en proyectos", style="Heading 1")
    doc.add_paragraph("Planificación, aseguramiento y control de la calidad",
                       style="Subtitle")
    doc.add_paragraph("La gestión de la calidad se organiza en tres procesos.")
    ruta = tmp_path / "modulo.docx"
    doc.save(ruta)

    marcadores = {"item_4": "1.4. Gestión de la calidad en proyectos"}
    secciones, _img, faltantes, _com, _otp = segmentar_docx(ruta, marcadores)

    assert "item_4" not in faltantes
    assert "<h3>Planificación, aseguramiento y control de la calidad</h3>" \
        in secciones.get("item_4", "")


def test_subrayado_se_preserva(tmp_path):
    """Mammoth no preserva el subrayado por defecto (lo trata como una
    elección de estilo sin significado semántico): un asesor que marca los
    títulos de un expander subrayándolos ("Para maquetación: expander
    (títulos subrayados)") los perdía del todo — quedaban como <p> sueltos,
    indistinguibles del resto, y _es_encabezado_de_seccion() nunca los
    reconocía como encabezado de sección."""
    doc = docx.Document()
    doc.add_paragraph("1.1. Una sección", style="Heading 1")
    p = doc.add_paragraph()
    run = p.add_run("Un título subrayado")
    run.underline = True
    doc.add_paragraph("Cuerpo del párrafo.")
    ruta = tmp_path / "modulo.docx"
    doc.save(ruta)

    marcadores = {"item_1": "1.1. Una sección"}
    secciones, _img, _faltantes, _com, _otp = segmentar_docx(ruta, marcadores)

    assert "<u>Un título subrayado</u>" in secciones.get("item_1", "")
